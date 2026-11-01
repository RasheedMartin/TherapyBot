"""Reading text out of uploaded files (PDF, Word, plain text). Scanned PDFs have no text layer and are refused."""

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_PDF_PAGES = 200
SUPPORTED = (".pdf", ".docx", ".txt", ".md")


class UnreadableFile(ValueError):
    """A file we can't get text from; the message is shown to the person who uploaded it."""


def extract_text(upload, max_bytes: int = MAX_UPLOAD_BYTES, max_pages: int = MAX_PDF_PAGES) -> str:
    name = (upload.name or "").lower()
    if upload.size > max_bytes:
        raise UnreadableFile(f"That file is too large. Please upload a file under {max_bytes // (1024 * 1024)} MB.")
    if name.endswith(".pdf"):
        text = _pdf_text(upload, max_pages)
    elif name.endswith(".docx"):
        text = _docx_text(upload)
    elif name.endswith((".txt", ".md")):
        text = upload.read().decode("utf-8", errors="replace")
    elif name.endswith(".doc"):
        raise UnreadableFile("Older .doc files aren't supported. Save it as .docx or PDF and upload that.")
    else:
        raise UnreadableFile("Upload a PDF, Word (.docx) or text file.")
    return text.replace("\r\n", "\n").replace("\x00", "")


def _pdf_text(upload, max_pages: int) -> str:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(upload)
        text = "\n".join((page.extract_text() or "") for page in reader.pages[:max_pages])
    except (PdfReadError, ValueError) as e:
        raise UnreadableFile("Couldn't read that PDF. If it's password-protected, remove the password first.") from e
    if len(text.strip()) < 20:
        raise UnreadableFile("That PDF looks like a scanned image, so there's no text to read. "
                             "Export it as text-based PDF or Word from the original system instead.")
    return text


def _docx_text(upload) -> str:
    from docx import Document
    from docx.opc.exceptions import PackageNotFoundError

    try:
        document = Document(upload)
    except (PackageNotFoundError, ValueError, KeyError) as e:
        raise UnreadableFile("Couldn't read that Word file. Try saving it again as .docx or PDF.") from e
    lines = [p.text for p in document.paragraphs]
    for table in document.tables:  # notes exported as forms often keep fields in tables
        for row in table.rows:
            cells = list(dict.fromkeys(cell.text.strip() for cell in row.cells if cell.text.strip()))
            if cells:
                lines.append(" | ".join(cells))
    text = "\n".join(lines)
    if len(text.strip()) < 20:
        raise UnreadableFile("That Word file doesn't contain any text.")
    return text
